Import argparse
import json
import os
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _shift_month(value: date, months: int) -> date:
    year, month = divmod(value.year * 12 + value.month - 1 + months, 12)
    return date(year, month + 1, 1)


def _get_report_period(today=None):
    this_month = (today or date.today()).replace(day=1)
    return _shift_month(this_month, -1), this_month


def _percentage_change(current, previous):
    if previous == 0:
        return None
    return ((current - previous) / previous * Decimal('100')).quantize(
        Decimal('0.01')
    )


def _validate_config(config):
    environments = config.get('environments', {})
    if not 1 <= len(environments) <= 2 or not set(environments).issubset({'Prod', 'Non-Prod'}):
        raise ValueError('Configure Prod, Non-Prod, or both.')
    seen = set()
    for name, settings in environments.items():
        accounts = settings.get('accounts', [])
        if not accounts or len(accounts) != len(set(accounts)):
            raise ValueError(f'{name}: configure unique account IDs.')
        for account in accounts:
            if not isinstance(account, str) or len(account) != 12 or not account.isdigit():
                raise ValueError('Account IDs must be 12-digit strings.')
            if account in seen:
                raise ValueError('An account cannot be in both environments.')
            seen.add(account)
        link = settings.get('dashboard_url', '')
        if link and not link.startswith('https://'):
            raise ValueError('Dashboard URLs must use HTTPS.')
    if config.get('metric') not in {
        'UnblendedCost', 'NetUnblendedCost', 'AmortizedCost', 'NetAmortizedCost'
    }:
        raise ValueError('Configure a supported cost metric.')
    currency = config.get('currency', '')
    if not isinstance(currency, str) or len(currency) != 3 or not currency.isalpha() or not currency.isupper():
        raise ValueError('Configure the source currency, for example USD.')
    if config.get('display_months', 3) not in (3, 4):
        raise ValueError('display_months must be 3 or 4.')


def _get_monthly_costs(client, start_date, end_date, accounts, metric, currency):
    request = {
        'TimePeriod': {'Start': start_date.isoformat(), 'End': end_date.isoformat()},
        'Granularity': 'MONTHLY',
        'Metrics': [metric],
        'Filter': {'Dimensions': {'Key': 'LINKED_ACCOUNT', 'Values': accounts}},
    }
    costs, estimated = {}, []
    while True:
        response = client.get_cost_and_usage(**request)
        for result in response.get('ResultsByTime', []):
            month = result['TimePeriod']['Start'][:7]
            value = result['Total'][metric]
            if value['Unit'] != currency:
                raise ValueError(f"Currency mismatch: {value['Unit']} vs {currency}.")
            if month in costs:
                raise ValueError(f'Duplicate monthly total: {month}.')
            costs[month] = Decimal(value['Amount'])
            if result.get('Estimated'):
                estimated.append(month)
        token = response.get('NextPageToken')
        if not token:
            return costs, estimated
        request['NextPageToken'] = token


def _build_report(costs, report_month, count=3):
    report_month = report_month.replace(day=1)
    months = [
        _shift_month(report_month, offset).strftime('%Y-%m')
        for offset in range(1 - count, 1)
    ]
    baseline = _shift_month(report_month, -count).strftime('%Y-%m')
    prior_year = _shift_month(report_month, -12).strftime('%Y-%m')
    required = set(months + [baseline, prior_year])
    rows = {}
    for environment in costs:
        amounts = costs[environment]
        missing = required - amounts.keys()
        if missing:
            raise ValueError(f'{environment}: missing periods {sorted(missing)}.')
        rows[environment] = {
            'costs': [amounts[month] for month in months],
            'prior_year_cost': amounts[prior_year],
            'yoy': _percentage_change(amounts[months[-1]], amounts[prior_year]),
            'mom': [
                _percentage_change(amounts[month], amounts[previous])
                for month, previous in zip(months, [baseline] + months[:-1])
            ],
        }
    totals = {key: sum(costs[name][key] for name in rows) for key in required}
    return {
        'report_month': report_month.strftime('%Y-%m'),
        'months': months,
        'environments': rows,
        'total_mom': [
            _percentage_change(totals[month], totals[previous])
            for month, previous in zip(months, [baseline] + months[:-1])
        ],
        'total_yoy': _percentage_change(totals[months[-1]], totals[prior_year]),
    }


def _save_report(report, config, output, demo=False):
    output.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Monthly costs'
    sheet.append(['SAMPLE DATA' if demo else 'AWS monthly costs', report['report_month']])
    sheet.append(['Currency', config['currency'], 'Metric', config['metric']])
    sheet.append(['Environment', *report['months'], 'Same month last year', 'YoY', 'Dashboard'])
    count = len(report['months'])
    as_percent = lambda value: None if value is None else float(value / 100)
    for environment, row in report['environments'].items():
        settings = config.get('environments', {}).get(environment, {})
        sheet.append([
            environment, *[float(cost) for cost in row['costs']],
            float(row['prior_year_cost']), as_percent(row['yoy']),
            settings.get('dashboard_url') or 'Pending',
        ])
        for column in range(2, count + 3):
            sheet.cell(sheet.max_row, column).number_format = '#,##0.00'
        sheet.cell(sheet.max_row, count + 3).number_format = '0.00%'
    sheet.append(['Total MoM', *map(as_percent, report['total_mom']), None, as_percent(report['total_yoy'])])
    for column in list(range(2, count + 2)) + [count + 3]:
        sheet.cell(sheet.max_row, column).number_format = '0.00%'
    sheet.append([])
    sheet.append(['Blank percentages mean N/A: the comparison cost was zero.'])
    sheet.append(['No currency conversion. Anomaly integration is pending.'])
    sheet.append(['Estimated cost periods', json.dumps(report.get('estimated_periods', {}))])
    for cell in sheet[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E8EFF7')
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    sheet.freeze_panes = 'B4'
    detail = workbook.create_sheet('Environment MoM')
    detail.append(['Environment', 'Month', 'MoM'])
    for environment, row in report['environments'].items():
        for month, change in zip(report['months'], row['mom']):
            detail.append([environment, month, as_percent(change)])
            detail.cell(detail.max_row, 3).number_format = '0.00%'
    workbook.save(output / 'monthly-cost-report.xlsx')
    (output / 'monthly-cost-report.json').write_text(
        json.dumps(report, default=str, indent=2), encoding='utf-8'
    )


def lambda_handler(event, context):
    """Generate a completed-month report using the Lambda execution role.

    REPORT_CONFIG_JSON uses the same schema as the local --config file,
    without a local AWS profile. REPORT_BUCKET must name an existing bucket.
    REPORT_PREFIX defaults to aws-cost-reports. An optional event field,
    report_month, selects a completed month in YYYY-MM format.
    """
    if not isinstance(event, dict):
        raise ValueError('The Lambda event must be a JSON object.')
    from .cost_report_notifications import notification_settings, publish_report_email

    email_settings = notification_settings(os.environ)
    raw_config = os.environ.get('REPORT_CONFIG_JSON')
    bucket = os.environ.get('REPORT_BUCKET', '').strip()
    if not raw_config:
        raise ValueError('Configure the REPORT_CONFIG_JSON environment variable.')
    if not bucket:
        raise ValueError('Configure the REPORT_BUCKET environment variable.')
    config = json.loads(raw_config)
    if not isinstance(config, dict):
        raise ValueError('REPORT_CONFIG_JSON must contain a JSON object.')
    deployment_env = os.environ.get('ENVIRONMENT', '').strip().lower()
    if deployment_env not in {'dev', 'test', 'prod'}:
        raise ValueError('ENVIRONMENT must be dev, test or prod.')

    environment_name = 'Prod' if deployment_env == 'prod' else 'Non-Prod'
    account_id = context.invoked_function_arn.split(':')[4]

    settings = dict(config.get('environments', {}).get(environment_name, {}))
    settings['accounts'] = [account_id]
    config['environments'] = {environment_name: settings}
    _validate_config(config)
    if config.get('profile'):
        raise ValueError('Remove profile from Lambda configuration; use its execution role.')

    requested_month = event.get('report_month')
    if requested_month is None:
        month = _get_report_period()[0]
    else:
        if not isinstance(requested_month, str):
            raise ValueError('report_month must use YYYY-MM format.')
        try:
            month = date.fromisoformat(requested_month + '-01')
        except ValueError as error:
            raise ValueError('report_month must use YYYY-MM format.') from error
        if month.strftime('%Y-%m') != requested_month:
            raise ValueError('report_month must use YYYY-MM format.')
    if month >= date.today().replace(day=1):
        raise ValueError('Select a completed month, before the current month.')

    import boto3
    from botocore.config import Config

    session = boto3.Session()
    client = session.client('ce', region_name='us-east-1', config=Config(
        retries={'mode': 'standard', 'max_attempts': 5}
    ))
    costs, estimated = {}, {}
    for name, settings in config['environments'].items():
        costs[name], estimated[name] = _get_monthly_costs(
            client, _shift_month(month, -12), _shift_month(month, 1),
            settings['accounts'], config['metric'], config['currency']
        )
    report = _build_report(costs, month, config.get('display_months', 3))
    report.update(currency=config['currency'], metric=config['metric'], estimated_periods=estimated)

    s3 = session.client('s3')
    prefix = os.environ.get('REPORT_PREFIX', 'aws-cost-reports').strip('/')
    folder = '/'.join(part for part in (prefix, report['report_month']) if part)
    artifacts = {}
    with TemporaryDirectory(prefix='aws-cost-report-', dir='/tmp') as directory:
        output = Path(directory)
        _save_report(report, config, output)
        for filename, content_type in (
            ('monthly-cost-report.xlsx',
             'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
            ('monthly-cost-report.json', 'application/json'),
        ):
            key = f'{folder}/{filename}'
            s3.upload_file(str(output / filename), bucket, key,
                           ExtraArgs={'ContentType': content_type})
            artifacts[filename] = {'bucket': bucket, 'key': key}
    notification = publish_report_email(
        session, email_settings, report, environment_name, account_id,
        artifacts['monthly-cost-report.xlsx'],
    )
    return {
        'notification': notification,
        'report_month': report['report_month'],
        'artifacts': artifacts,
        'estimated_periods': estimated,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument('--demo', action='store_true')
    source.add_argument('--config', type=Path)
    parser.add_argument('--month', help='Completed reporting month: YYYY-MM')
    parser.add_argument('--output', type=Path, default=Path('report-output'))
    args = parser.parse_args()
    month = date.fromisoformat(args.month + '-01') if args.month else _get_report_period()[0]
    if month >= date.today().replace(day=1):
        parser.error('Select a completed month, before the current month.')
    if args.demo:
        config = {'currency': 'USD', 'metric': 'UnblendedCost', 'display_months': 3}
        costs = {}
        for name, base in [('Prod', 1000), ('Non-Prod', 400)]:
            costs[name] = {
                _shift_month(month, offset).strftime('%Y-%m'): Decimal(base + (offset + 12) * 50)
                for offset in range(-12, 1)
            }
        estimated = {}
    else:
        import boto3
        from botocore.config import Config

        config = json.loads(args.config.read_text(encoding='utf-8'))
        _validate_config(config)
        session = boto3.Session(profile_name=config.get('profile'))
        client = session.client('ce', region_name='us-east-1', config=Config(
            retries={'mode': 'standard', 'max_attempts': 5}
        ))
        costs, estimated = {}, {}
        for name, settings in config['environments'].items():
            costs[name], estimated[name] = _get_monthly_costs(
                client, _shift_month(month, -12), _shift_month(month, 1),
                settings['accounts'], config['metric'], config['currency']
            )
    report = _build_report(costs, month, config.get('display_months', 3))
    report.update(currency=config['currency'], metric=config['metric'], estimated_periods=estimated)
    _save_report(report, config, args.output, args.demo)
    print(f'Report saved to: {args.output.resolve()}')


if __name__ == '__main__':
    main()